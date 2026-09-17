#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""测试用例 Excel 生成器（通用版）：python build_testcase_xlsx.py <用例数据.json> [输出.xlsx]
仅依赖 openpyxl（缺失: pip install openpyxl）。"""

import json
import math
import os
import sys


def display_width(text):
    """估算显示宽度：中文字符按 2 计，ASCII 按 1 计（用于行高估算）。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def main():
    if len(sys.argv) < 2:
        print("用法: python build_testcase_xlsx.py <用例数据.json> [输出文件.xlsx]")
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.isfile(input_path):
        print(f"[错误] 找不到输入文件: {input_path}")
        sys.exit(1)

    with open(input_path, encoding="utf-8") as f:
        spec = json.load(f)

    default_name = spec.get("file_name") or os.path.splitext(os.path.basename(input_path))[0]
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(input_path)), default_name + ".xlsx"
    )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("[错误] 缺少 openpyxl 库，请先执行: pip install openpyxl")
        sys.exit(1)

    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(name="微软雅黑", size=10)
    data_align = Alignment(wrap_text=True, vertical="top", horizontal="left")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    wb.remove(wb.active)  # 删除默认空 Sheet

    total_rows = 0
    for sheet in spec.get("sheets", []):
        headers = sheet["headers"]
        widths = sheet.get("col_widths") or [16] * len(headers)
        rows = sheet.get("rows", [])
        if len(headers) != len(widths):
            raise ValueError(f"Sheet[{sheet.get('name')}] headers 与 col_widths 长度不一致")

        ws = wb.create_sheet(sheet.get("name", "Sheet1")[:31])

        for col, (name, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=col, value=str(name))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 24

        for r, row in enumerate(rows, start=2):
            if len(row) != len(headers):
                raise ValueError(
                    f"Sheet[{sheet.get('name')}] 第 {r - 1} 条数据列数({len(row)})"
                    f"与表头({len(headers)})不一致"
                )
            max_lines = 1
            for col, value in enumerate(row, start=1):
                text = "" if value is None else str(value)
                cell = ws.cell(row=r, column=col, value=text)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = border
                usable = max(widths[col - 1] - 2, 4)
                lines = sum(max(1, math.ceil(display_width(seg) / usable)) for seg in text.split("\n"))
                max_lines = max(max_lines, lines)
            ws.row_dimensions[r].height = max(18, min(16 * max_lines + 8, 320))

        total_rows += len(rows)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    print(f"已生成: {out_path}")
    print(f"汇总: {len(spec.get('sheets', []))} 个 Sheet，共 {total_rows} 条用例")


if __name__ == "__main__":
    main()
