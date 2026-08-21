#!/usr/bin/env python3
"""Read-only source extraction for WCS requirement analysis.

The script emits JSON Lines. It preserves format-specific locators but does not
decide whether a statement belongs to WCS. Visual review remains mandatory for
layout-sensitive, diagram-heavy, merged, scanned, or otherwise low-text files.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable, Iterator


SUPPORTED = {".docx", ".xlsx", ".xlsm", ".pdf", ".txt", ".md", ".csv", ".tsv"}
LEGACY_OR_UNSUPPORTED = {".doc", ".xls", ".rtf", ".wps", ".et"}
TEMP_PREFIXES = ("~$", ".~lock.")
WCS_ANCHOR_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])WCS(?![A-Za-z0-9])|Warehouse\s+Control\s+System|仓储控制系统|仓库控制系统",
    re.IGNORECASE,
)
POTENTIAL_WCS_ALIAS_PATTERN = re.compile(r"调度系统|仓储调度|设备控制系统|物流控制系统|控制系统")
BOUNDARY_SYSTEM_PATTERN = re.compile(
    r"\b(?:WMS|WES|MES|ERP|OMS|NCC|PLC|SCADA|RCS|AGV|AMR)\b|上游系统|下游系统|设备控制器|机器人控制系统|设备厂商",
    re.IGNORECASE,
)
VAGUE_SCOPE_PATTERN = re.compile(
    r"支持|具备|实现|提供|相关|等功能|智能(?:化)?|自动(?:化)?|优化|灵活|高效|稳定|完善|无缝|兼容|适配|及时|按需|必要时|可配置"
)
GENERIC_ACTOR_PATTERN = re.compile(r"(?:^|[，。；：\s])(系统|平台|软件|上位系统|控制端)(?:应|需|可|支持|负责|实现|提供)")
INTERFACE_PATTERN = re.compile(r"接口|对接|交互|通信|报文|消息|API|MQ|回传|下发", re.IGNORECASE)
EXCEPTION_PATTERN = re.compile(r"异常|故障|超时|离线|重试|恢复|补偿|人工接管|手动|强制完成|断点续传|死信")
DATA_OWNERSHIP_PATTERN = re.compile(r"库存|储位|库位|任务状态|设备状态|主数据|基础数据|物料档案|用户权限|缓存|同步|回写")
ACCEPTANCE_GAP_PATTERN = re.compile(r"实时|高效|稳定|灵活|兼容|无缝|自动|智能|可靠|快速|友好")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def add_analysis_hints(record: dict[str, Any]) -> dict[str, Any]:
    """Add deterministic triage hints without making a final WCS classification."""
    if record.get("record_type") != "content":
        return record
    text = clean_text(record.get("text"))
    if not text:
        return record

    explicit_wcs = bool(WCS_ANCHOR_PATTERN.search(text))
    potential_alias = bool(POTENTIAL_WCS_ALIAS_PATTERN.search(text)) and not explicit_wcs
    signals: list[str] = []
    if BOUNDARY_SYSTEM_PATTERN.search(text):
        signals.append("cross_system_boundary")
    if VAGUE_SCOPE_PATTERN.search(text):
        signals.append("vague_or_capability_language")
    if GENERIC_ACTOR_PATTERN.search(text) and not explicit_wcs:
        signals.append("generic_actor_reference")
    if INTERFACE_PATTERN.search(text):
        signals.append("interface_contract_review")
    if EXCEPTION_PATTERN.search(text):
        signals.append("exception_ownership_review")
    if DATA_OWNERSHIP_PATTERN.search(text):
        signals.append("data_authority_review")
    if ACCEPTANCE_GAP_PATTERN.search(text):
        signals.append("acceptance_measure_review")

    if explicit_wcs or potential_alias or signals:
        record["analysis_hints"] = {
            "explicit_wcs_anchor": explicit_wcs,
            "potential_wcs_alias": potential_alias,
            "priority_review_candidate": bool((explicit_wcs or potential_alias) and signals),
            "signals": signals,
            "note": "Triage hints only; verify context, responsibility, ambiguity, and emphasis level using the skill rules.",
        }
    return record


def clipped(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars] + "…", True


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def base_record(path: Path, source_id: str) -> dict[str, Any]:
    return {
        "source_id": source_id,
        "file": path.name,
        "path": str(path.resolve()),
        "format": path.suffix.lower().lstrip("."),
    }


def manifest_record(path: Path, source_id: str) -> dict[str, Any]:
    record = base_record(path, source_id)
    record.update(
        {
            "record_type": "source_manifest",
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
            "status": "extracting" if path.suffix.lower() in SUPPORTED else "conversion_required",
        }
    )
    return record


def dependency_error(path: Path, source_id: str, package: str, exc: Exception) -> dict[str, Any]:
    record = base_record(path, source_id)
    record.update(
        {
            "record_type": "error",
            "status": "dependency_missing",
            "required_package": package,
            "message": str(exc),
        }
    )
    return record


def extract_docx(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        yield dependency_error(path, source_id, "python-docx", exc)
        return

    document = Document(path)
    headings: list[tuple[int, str]] = []
    heading_pattern = re.compile(r"(?:Heading|标题)\s*([1-9])", re.IGNORECASE)

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = clean_text(paragraph.text)
        if not text:
            continue
        style_name = clean_text(getattr(paragraph.style, "name", ""))
        match = heading_pattern.search(style_name)
        if match:
            level = int(match.group(1))
            headings = [(existing_level, heading) for existing_level, heading in headings if existing_level < level]
            headings.append((level, text))
        body, was_clipped = clipped(text, max_chars)
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "content",
                "locator": f"paragraph {index}",
                "location_kind": "docx_paragraph",
                "text": body,
                "metadata": {
                    "paragraph_index": index,
                    "style": style_name,
                    "heading_path": [heading for _, heading in headings],
                    "truncated": was_clipped,
                },
            }
        )
        yield record

    for table_index, table in enumerate(document.tables, start=1):
        headers = [clean_text(cell.text) for cell in table.rows[0].cells] if table.rows else []
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                text = clean_text(cell.text)
                if not text:
                    continue
                body, was_clipped = clipped(text, max_chars)
                header = headers[column_index - 1] if column_index <= len(headers) else ""
                record = base_record(path, source_id)
                record.update(
                    {
                        "record_type": "content",
                        "locator": f"table {table_index} / row {row_index} / column {column_index}",
                        "location_kind": "docx_table_cell",
                        "text": body,
                        "metadata": {
                            "table_index": table_index,
                            "row_index": row_index,
                            "column_index": column_index,
                            "column_header": header,
                            "truncated": was_clipped,
                        },
                    }
                )
                yield record

    seen_header_footer: set[tuple[str, int, int, str]] = set()
    for section_index, section in enumerate(document.sections, start=1):
        for container_name, container in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(container.paragraphs, start=1):
                text = clean_text(paragraph.text)
                if not text:
                    continue
                dedupe_key = (container_name, paragraph_index, len(text), text)
                if dedupe_key in seen_header_footer:
                    continue
                seen_header_footer.add(dedupe_key)
                body, was_clipped = clipped(text, max_chars)
                record = base_record(path, source_id)
                record.update(
                    {
                        "record_type": "content",
                        "locator": f"section {section_index} / {container_name} / paragraph {paragraph_index}",
                        "location_kind": f"docx_{container_name}",
                        "text": body,
                        "metadata": {"truncated": was_clipped},
                    }
                )
                yield record

    warning = base_record(path, source_id)
    warning.update(
        {
            "record_type": "warning",
            "status": "visual_review_required",
            "message": "DOCX extraction may omit text boxes, drawings, comments, and tracked-change semantics; render and inspect the document before final citation.",
        }
    )
    yield warning


def merged_range_for(sheet: Any, coordinate: str) -> str | None:
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            return str(merged_range)
    return None


def extract_xlsx(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        yield dependency_error(path, source_id, "openpyxl", exc)
        return

    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, data_only=False, read_only=False, keep_vba=keep_vba)
    cached_workbook = load_workbook(path, data_only=True, read_only=False, keep_vba=keep_vba)
    try:
        for sheet in workbook.worksheets:
            cached_sheet = cached_workbook[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None and cell.comment is None:
                        continue
                    formula = cell.value if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")) else None
                    cached_value = cached_sheet[cell.coordinate].value
                    visible_value = cached_value if formula is not None and cached_value is not None else cell.value
                    text = clean_text(visible_value)
                    if not text and formula is not None:
                        text = clean_text(formula)
                    body, was_clipped = clipped(text, max_chars)
                    merged_range = merged_range_for(sheet, cell.coordinate)
                    row_hidden = bool(sheet.row_dimensions[cell.row].hidden)
                    column_hidden = bool(sheet.column_dimensions[cell.column_letter].hidden)
                    record = base_record(path, source_id)
                    record.update(
                        {
                            "record_type": "content",
                            "locator": f"{sheet.title}!{merged_range or cell.coordinate}",
                            "location_kind": "xlsx_cell",
                            "text": body,
                            "metadata": {
                                "sheet": sheet.title,
                                "sheet_state": sheet.sheet_state,
                                "cell": cell.coordinate,
                                "merged_range": merged_range,
                                "formula": clean_text(formula),
                                "cached_value": clean_text(cached_value),
                                "comment": clean_text(cell.comment.text) if cell.comment else "",
                                "row_hidden": row_hidden,
                                "column_hidden": column_hidden,
                                "truncated": was_clipped,
                            },
                        }
                    )
                    yield record
    finally:
        workbook.close()
        cached_workbook.close()

    warning = base_record(path, source_id)
    warning.update(
        {
            "record_type": "warning",
            "status": "visual_review_required",
            "message": "XLSX extraction does not evaluate formulas and may miss requirements encoded only by formatting, drawings, charts, controls, or layout; render every relevant sheet before final citation.",
        }
    )
    yield warning


def extract_pdf(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        yield dependency_error(path, source_id, "pypdf", exc)
        return

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            result = reader.decrypt("")
        except Exception as exc:  # library-specific exceptions vary
            result = 0
            decrypt_error = str(exc)
        else:
            decrypt_error = ""
        if not result:
            record = base_record(path, source_id)
            record.update(
                {
                    "record_type": "error",
                    "status": "protected",
                    "message": decrypt_error or "PDF requires a password.",
                }
            )
            yield record
            return

    for page_number, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        body, was_clipped = clipped(text, max_chars)
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "content",
                "locator": f"page {page_number}",
                "location_kind": "pdf_page",
                "text": body,
                "metadata": {
                    "page_number": page_number,
                    "character_count": len(text),
                    "needs_visual_or_ocr_review": len(text) < 20,
                    "truncated": was_clipped,
                },
            }
        )
        yield record


def read_decoded_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replacement"


def extract_text_file(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    content, encoding = read_decoded_text(path)
    for line_number, line in enumerate(content.splitlines(), start=1):
        text = clean_text(line)
        if not text:
            continue
        body, was_clipped = clipped(text, max_chars)
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "content",
                "locator": f"line {line_number}",
                "location_kind": "text_line",
                "text": body,
                "metadata": {"line_number": line_number, "encoding": encoding, "truncated": was_clipped},
            }
        )
        yield record


def extract_delimited(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    content, encoding = read_decoded_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows = list(csv.reader(content.splitlines(), delimiter=delimiter))
    headers = rows[0] if rows else []
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            text = clean_text(value)
            if not text:
                continue
            body, was_clipped = clipped(text, max_chars)
            header = clean_text(headers[column_index - 1]) if column_index <= len(headers) else ""
            record = base_record(path, source_id)
            record.update(
                {
                    "record_type": "content",
                    "locator": f"row {row_index} / column {column_index}",
                    "location_kind": "delimited_cell",
                    "text": body,
                    "metadata": {
                        "row_index": row_index,
                        "column_index": column_index,
                        "column_header": header,
                        "encoding": encoding,
                        "truncated": was_clipped,
                    },
                }
            )
            yield record


def iter_input_files(inputs: Iterable[str], recursive: bool) -> list[Path]:
    files: list[Path] = []
    for raw_input in inputs:
        path = Path(raw_input).expanduser()
        if path.is_file():
            files.append(path)
        elif path.is_dir():
            iterator = path.rglob("*") if recursive else path.glob("*")
            files.extend(candidate for candidate in iterator if candidate.is_file())
        else:
            files.append(path)
    unique: dict[str, Path] = {}
    for path in files:
        key = str(path.resolve()).casefold() if path.exists() else str(path.absolute()).casefold()
        unique[key] = path
    return sorted(unique.values(), key=lambda item: str(item).casefold())


def extract_file(path: Path, source_id: str, max_chars: int) -> Iterator[dict[str, Any]]:
    if not path.exists():
        record = base_record(path, source_id)
        record.update({"record_type": "error", "status": "missing", "message": "Input path does not exist."})
        yield record
        return
    if path.name.startswith(TEMP_PREFIXES):
        record = base_record(path, source_id)
        record.update({"record_type": "warning", "status": "skipped_temp_file", "message": "Temporary Office file skipped."})
        yield record
        return

    yield manifest_record(path, source_id)
    suffix = path.suffix.lower()
    if suffix in LEGACY_OR_UNSUPPORTED:
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "warning",
                "status": "conversion_required",
                "message": "Convert to a reviewable modern format and preserve the original for version traceability.",
            }
        )
        yield record
        return
    if suffix not in SUPPORTED:
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "warning",
                "status": "unsupported",
                "message": "Use the relevant format-specific workflow and add a manual source locator.",
            }
        )
        yield record
        return

    extractors = {
        ".docx": extract_docx,
        ".xlsx": extract_xlsx,
        ".xlsm": extract_xlsx,
        ".pdf": extract_pdf,
        ".txt": extract_text_file,
        ".md": extract_text_file,
        ".csv": extract_delimited,
        ".tsv": extract_delimited,
    }
    try:
        yield from extractors[suffix](path, source_id, max_chars)
    except Exception as exc:  # continue the corpus and expose the failed file
        record = base_record(path, source_id)
        record.update(
            {
                "record_type": "error",
                "status": "extraction_failed",
                "error_type": type(exc).__name__,
                "message": str(exc),
            }
        )
        yield record


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract source text and stable locators as JSONL without modifying input files.")
    parser.add_argument("paths", nargs="+", help="Source files or directories")
    parser.add_argument("--recursive", action="store_true", help="Recurse into input directories")
    parser.add_argument("--output", help="Write JSONL to this file; stdout when omitted")
    parser.add_argument("--max-chars", type=int, default=12000, help="Maximum text characters per content record")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.max_chars < 200:
        raise SystemExit("--max-chars must be at least 200")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    files = iter_input_files(args.paths, args.recursive)
    output_stream = None
    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_stream = output_path.open("w", encoding="utf-8", newline="\n")
    stream = output_stream or sys.stdout

    error_count = 0
    try:
        for source_number, path in enumerate(files, start=1):
            source_id = f"SRC-{source_number:03d}"
            for record in extract_file(path, source_id, args.max_chars):
                record = add_analysis_hints(record)
                if record.get("record_type") == "error":
                    error_count += 1
                stream.write(json.dumps(record, ensure_ascii=False) + "\n")
    finally:
        if output_stream is not None:
            output_stream.close()

    return 1 if error_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
